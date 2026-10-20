#!/usr/bin/env python3
"""Re-run protection: captioning a deck that already carries captioner captions must not duplicate them.

Runs the real extract -> apply chain on the bundled deck (examples/sample.pptx), then runs captioner
again on its own output:

  1. first run writes captions (positive control: the output holds captioner shapes)
  2. a second run WITHOUT --update-existing refuses: exit 4, the message names the flag, and it writes
     neither a deck nor an audit CSV for that deck
  3. a second run WITH --update-existing succeeds and the deck holds the same number of captioner
     shapes as after the first run (replaced, not duplicated)
  4. --update-existing also strips in-picture band captions (captioner_capband_), which the bundled
     deck does not produce, so this one is a unit check on an in-memory slide

Deterministic, no vision call: captions come from the committed examples/sample_captions.json.
Run:  python3 tests/test_rerun.py     (exit 0 = all pass, 1 = a failure)
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
import apply_captions as ac  # noqa: E402
SAMPLE = REPO / "examples" / "sample.pptx"
CAPTIONS_FIXTURE = REPO / "examples" / "sample_captions.json"
PREFIXES = ("captioner_caption_", "captioner_capband_", "captioner_smartart_", "captioner_sa_icon_")
FAILS = []


def check(name, cond):
    print(f"  {'PASS' if cond else 'FAIL'}  {name}")
    if not cond:
        FAILS.append(name)


def captioner_shapes(path: Path) -> int:
    n = 0
    for slide in Presentation(str(path)).slides:
        for sh in slide.shapes:
            if (sh.name or "").startswith(PREFIXES):
                n += 1
    return n


def extract(deck: Path, work: Path):
    subprocess.run([sys.executable, str(REPO / "scripts" / "extract_images.py"), str(deck), str(work)],
                   check=True, capture_output=True, text=True)


def apply(work: Path, *flags):
    return subprocess.run([sys.executable, str(REPO / "scripts" / "apply_captions.py"), str(work),
                           "--no-spellcheck", *flags], capture_output=True, text=True)


def captions_for(deck_name: str) -> dict:
    fixture = json.loads(CAPTIONS_FIXTURE.read_text())
    return {f"{deck_name}/{k.split('/', 1)[1]}": v for k, v in fixture.items()}


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="caprerun_"))
    try:
        # 1. first run on the original deck
        w1 = tmp / "first"
        extract(SAMPLE, w1)
        (w1 / "captions.json").write_text(json.dumps(captions_for("sample")))
        r1 = apply(w1)
        out1 = w1 / "captioned_decks" / "sample_captioned.pptx"
        n1 = captioner_shapes(out1) if out1.exists() else 0
        check("first run writes captions (positive control)", r1.returncode == 0 and n1 > 0)

        # 2. second run on the captioned output, without --update-existing
        again = tmp / "sample_captioned.pptx"
        shutil.copy(out1, again)
        w2 = tmp / "second"
        extract(again, w2)
        (w2 / "captions.json").write_text(json.dumps(captions_for("sample_captioned")))
        r2 = apply(w2)
        out2 = w2 / "captioned_decks" / "sample_captioned.pptx"
        check("a second run without --update-existing refuses (exit 4)", r2.returncode == 4)
        check("the refusal names --update-existing", "--update-existing" in r2.stdout)
        check("a refused run writes no deck", not out2.exists())
        check("a refused run writes no audit CSV for that deck",
              not (w2 / "audit" / "sample_captioned_audit.csv").exists())

        # 3. the same second run with --update-existing replaces instead of duplicating
        r3 = apply(w2, "--update-existing")
        n3 = captioner_shapes(out2) if out2.exists() else -1
        check("with --update-existing the re-run succeeds", r3.returncode == 0 and out2.exists())
        check("with --update-existing captions are replaced, not duplicated", n3 == n1)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # 4. band captions are captioner shapes too: counted, and stripped by --update-existing
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name in ("captioner_capband_abc123def456", "captioner_caption_0123456789ab", "a person's text box"):
        tb = slide.shapes.add_textbox(Emu(100000), Emu(100000), Emu(2000000), Emu(300000))
        tb.text_frame.text = "text"
        tb.name = name
    check("a band caption counts as a captioner shape", ac.count_captioner_shapes(prs) == 2)
    removed, _ = ac.remove_previous_caption_shapes(slide)
    left = [sh.name for sh in slide.shapes]
    check("--update-existing strips band captions as well as ordinary ones",
          removed == 2 and left == ["a person's text box"])

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(100000), Emu(100000), Emu(2000000), Emu(300000))
    tb.text_frame.text = "Photo: edited by a person."
    tb.name = "captioner_capband_abc123def456"
    removed, kept = ac.remove_previous_caption_shapes(slide, oracle={"abc123def456": "Photo: the original."})
    check("a band caption a person edited is kept", removed == 0 and kept == 1 and len(slide.shapes) == 1)

    print(f"test_rerun: {len(FAILS)} failed" if FAILS else "test_rerun: all passed")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
