"""Tests for the v0.3 WCAG alt-text (descr) extension.

Run directly:  python3 test_descr_writer.py
(also pytest-compatible: every test_* function is self-contained)

Covers: the pure decide_descr() decision table, prefix-strip, the narrow
autogen pattern list, decorative-marking XML round-trip (spike-2 verified
element), end-to-end apply (authored / preserved / decorative), idempotent
re-run via the oracle, hand-edit preservation, --replace-autogen-descr,
--descr-only, and the verify.py --descr coverage gate exit code.
"""
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import hashlib

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

# Resolve the scripts dir whether this test lives alongside the scripts (flat
# layout) or in a sibling tests/ dir (repo layout): prefer <here>/../scripts,
# else fall back to the test's own directory.
_HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(_HERE), 'scripts')
if not os.path.exists(os.path.join(SCRIPTS, '_descr_writer.py')):
    SCRIPTS = _HERE
sys.path.insert(0, SCRIPTS)

from _descr_writer import (          # noqa: E402
    decide_descr, derive_alt_text, is_autogen_descr, mark_decorative,
    is_decorative_marked, iter_live_pic_cnvprs, DECORATIVE_MARK,
    DECOR_EXT_URI, ADEC_NS,
)
from _oxml_pics import iter_slide_pics  # noqa: E402


def _png(color):
    buf = io.BytesIO()
    Image.new('RGB', (8, 8), color).save(buf, format='PNG')
    return buf.getvalue()


PNG_RED = _png((200, 30, 30))
PNG_GREEN = _png((30, 200, 30))
PNG_BLUE = _png((30, 30, 200))


def _h(blob):
    return hashlib.sha256(blob).hexdigest()[:12]


# ---------------------------------------------------------------- pure tests

def test_decide_descr_matrix():
    D = decide_descr
    # empty current
    assert D('', None, False, 'a mug.', False) == ('descr-authored', 'a mug.')
    assert D('', None, True, None, False) == ('descr-decorative-marked', DECORATIVE_MARK)
    assert D('', None, True, None, True) == ('descr-decorative-unchanged', None)
    assert D('', None, False, None, False) == ('descr-skipped-no-caption', None)
    # ours (oracle match) -> update/unchanged/reclassify
    assert D('a mug.', 'a mug.', False, 'a mug.', False) == ('descr-unchanged', None)
    assert D('a mug.', 'a mug.', False, 'a cup.', False) == ('descr-updated', 'a cup.')
    assert D('a mug.', 'a mug.', True, None, False) == ('descr-decorative-marked', DECORATIVE_MARK)
    # hand-edit (oracle mismatch) -> preserved unless forced
    assert D('edited by human', 'a mug.', False, 'a cup.', False) == \
        ('descr-preserved-hand-edit', None)
    assert D('edited by human', 'a mug.', False, 'a cup.', False,
             force_overwrite_edits=True) == ('descr-authored-forced', 'a cup.')
    # pre-existing unknown provenance -> preserved by default
    assert D('Hand authored text', None, False, 'a cup.', False) == \
        ('descr-preserved-pre-existing', None)
    # ... even when it LOOKS autogen, without the flag
    assert D('image.png', None, False, 'a cup.', False) == \
        ('descr-preserved-pre-existing', None)
    # autogen + flag -> replaced; non-autogen + flag -> still preserved
    assert D('image.png', None, False, 'a cup.', False, replace_autogen=True) == \
        ('descr-replaced-autogen', 'a cup.')
    assert D('Hand authored text', None, False, 'a cup.', False,
             replace_autogen=True) == ('descr-preserved-pre-existing', None)
    print('  decide_descr matrix OK')


def test_derive_alt_text():
    assert derive_alt_text('Photo: coffee mug on red background.') == \
        'coffee mug on red background.'
    assert derive_alt_text('Chart: revenue by quarter.') == 'revenue by quarter.'
    assert derive_alt_text('coffee mug.') == 'coffee mug.'   # no prefix -> verbatim
    assert derive_alt_text('Photo:') == 'Photo:'             # stripping would empty
    assert derive_alt_text(None) is None
    print('  derive_alt_text OK')


def test_is_autogen_descr():
    assert is_autogen_descr('A picture containing text, whiteboard '
                            'Description automatically generated')
    assert is_autogen_descr('Description automatically generated')
    assert is_autogen_descr('image.png')
    assert is_autogen_descr('Picture 1.JPG')
    assert is_autogen_descr('A picture containing food')
    assert not is_autogen_descr('Students working in the lab.')
    assert not is_autogen_descr('')
    print('  is_autogen_descr OK')


# ------------------------------------------------------------- deck helpers

def _build_deck(path, pics):
    """pics: list of (png_bytes, pre_existing_descr_or_None)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (png, descr) in enumerate(pics):
        shp = slide.shapes.add_picture(io.BytesIO(png), Inches(0.5 + 3.2 * i),
                                       Inches(1), Inches(2), Inches(2))
        cnvpr = shp._element.nvPicPr.cNvPr
        if descr is None:
            if 'descr' in cnvpr.attrib:      # python-pptx stamps the filename
                del cnvpr.attrib['descr']
        else:
            cnvpr.set('descr', descr)
    prs.save(path)


def _make_work(tmp, deck_path, captions_by_hash):
    """Minimal extract-style work dir for apply_captions.py."""
    work = os.path.join(tmp, 'work')
    os.makedirs(os.path.join(work, 'audit'), exist_ok=True)
    deck = os.path.splitext(os.path.basename(deck_path))[0]
    manifest = {deck: {'deck': deck, 'deck_path': deck_path,
                       'pictures': [{'image_hash': h, 'slide': 1}
                                    for h in captions_by_hash]}}
    with open(os.path.join(work, 'manifest.json'), 'w') as f:
        json.dump(manifest, f)
    with open(os.path.join(work, 'captions.json'), 'w') as f:
        json.dump({f'{deck}/{h}.png': c for h, c in captions_by_hash.items()}, f)
    return work, deck


def _apply(work, *extra):
    cmd = [sys.executable, os.path.join(SCRIPTS, 'apply_captions.py'), work,
           '--quick', '--write-descr', *extra]
    r = subprocess.run(cmd, capture_output=True, text=True)
    assert r.returncode == 0, f'apply failed rc={r.returncode}\n{r.stdout}\n{r.stderr}'
    return r.stdout


def _saved_slide_xml(work, deck):
    out = os.path.join(work, 'captioned_decks', f'{deck}_captioned.pptx')
    with zipfile.ZipFile(out) as z:
        return out, z.read('ppt/slides/slide1.xml').decode('utf-8')


def _audit_actions(work, deck):
    import csv as _csv
    p = os.path.join(work, 'audit', f'{deck}_descr_audit.csv')
    with open(p, newline='') as f:
        return {r['image_hash']: r['action'] for r in _csv.DictReader(f)}


# ------------------------------------------------------------ element tests

def test_mark_decorative_roundtrip():
    tmp = tempfile.mkdtemp(prefix='descr_t1_')
    try:
        deck_path = os.path.join(tmp, 'd.pptx')
        _build_deck(deck_path, [(PNG_RED, None)])
        prs = Presentation(deck_path)
        (el, cnvpr, pid) = iter_live_pic_cnvprs(prs.slides[0])[0]
        assert not is_decorative_marked(cnvpr)
        mark_decorative(cnvpr)
        mark_decorative(cnvpr)          # idempotent — no duplicate ext
        assert is_decorative_marked(cnvpr)
        prs.save(deck_path)
        with zipfile.ZipFile(deck_path) as z:
            xml = z.read('ppt/slides/slide1.xml').decode('utf-8')
        assert xml.count(DECOR_EXT_URI) == 1, 'duplicate decorative ext'
        assert ADEC_NS in xml and 'val="1"' in xml
        # reopen: detached read agrees
        prs2 = Presentation(deck_path)
        pic = iter_slide_pics(prs2.slides[0])[0]
        from _oxml_pics import NS as _NS
        assert is_decorative_marked(pic['el'].find('.//p:cNvPr', _NS))
        print('  mark_decorative round-trip OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------- end-to-end tests

def test_apply_end_to_end():
    tmp = tempfile.mkdtemp(prefix='descr_t2_')
    try:
        deck_path = os.path.join(tmp, 'deckA.pptx')
        _build_deck(deck_path, [(PNG_RED, None),
                                (PNG_GREEN, 'Hand authored text'),
                                (PNG_BLUE, None)])
        caps = {_h(PNG_RED): 'Photo: coffee mug on red background.',
                _h(PNG_GREEN): 'Photo: green square artwork.',
                _h(PNG_BLUE): '[decorative]'}
        work, deck = _make_work(tmp, deck_path, caps)
        _apply(work)
        out, xml = _saved_slide_xml(work, deck)
        # authored: prefix-stripped caption persisted into the SAVED file
        assert 'descr="coffee mug on red background."' in xml
        # preserved: human descr untouched
        assert 'descr="Hand authored text"' in xml
        # decorative: verified extension present, exactly once
        assert xml.count(DECOR_EXT_URI) == 1 and ADEC_NS in xml
        acts = _audit_actions(work, deck)
        assert acts[_h(PNG_RED)] == 'descr-authored'
        assert acts[_h(PNG_GREEN)] == 'descr-preserved-pre-existing'
        assert acts[_h(PNG_BLUE)] == 'descr-decorative-marked'

        # ---- idempotent re-run ON THE OUTPUT (oracle carried over) ----
        work2 = os.path.join(tmp, 'work2')
        os.makedirs(os.path.join(work2, 'audit'))
        shutil.copy(os.path.join(work, 'audit', f'{deck}_descr_audit.csv'),
                    os.path.join(work2, 'audit', f'{deck}_descr_audit.csv'))
        m2 = {deck: {'deck': deck, 'deck_path': out,
                     'pictures': [{'image_hash': h, 'slide': 1} for h in caps]}}
        with open(os.path.join(work2, 'manifest.json'), 'w') as f:
            json.dump(m2, f)
        with open(os.path.join(work2, 'captions.json'), 'w') as f:
            json.dump({f'{deck}/{h}.png': c for h, c in caps.items()}, f)
        subprocess.run([sys.executable, os.path.join(SCRIPTS, 'apply_captions.py'),
                        work2, '--quick', '--write-descr', '--update-existing'],
                       capture_output=True, text=True, check=True)
        acts2 = _audit_actions(work2, deck)
        assert acts2[_h(PNG_RED)] == 'descr-unchanged'
        assert acts2[_h(PNG_GREEN)] == 'descr-preserved-pre-existing'
        assert acts2[_h(PNG_BLUE)] == 'descr-decorative-unchanged'

        # ---- hand-edit on the output -> preserved on the NEXT run ----
        prs = Presentation(out)
        for (el, cnvpr, pid) in iter_live_pic_cnvprs(prs.slides[0]):
            if (cnvpr.get('descr') or '') == 'coffee mug on red background.':
                cnvpr.set('descr', 'my own better description')
        prs.save(out)
        subprocess.run([sys.executable, os.path.join(SCRIPTS, 'apply_captions.py'),
                        work2, '--quick', '--write-descr', '--update-existing'],
                       capture_output=True, text=True, check=True)
        acts3 = _audit_actions(work2, deck)
        assert acts3[_h(PNG_RED)] == 'descr-preserved-hand-edit'
        _, xml3 = _saved_slide_xml(work2, deck)
        assert 'descr="my own better description"' in xml3
        print('  end-to-end apply + idempotent re-run + hand-edit OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_replace_autogen_flag():
    tmp = tempfile.mkdtemp(prefix='descr_t3_')
    try:
        autogen = ('A picture containing text, whiteboard '
                   'Description automatically generated')
        deck_path = os.path.join(tmp, 'deckB.pptx')
        _build_deck(deck_path, [(PNG_RED, autogen), (PNG_GREEN, 'Real alt text.')])
        caps = {_h(PNG_RED): 'Photo: whiteboard sketch.',
                _h(PNG_GREEN): 'Photo: green square artwork.'}
        work, deck = _make_work(tmp, deck_path, caps)
        # default: BOTH preserved
        _apply(work)
        acts = _audit_actions(work, deck)
        assert acts[_h(PNG_RED)] == 'descr-preserved-pre-existing'
        assert acts[_h(PNG_GREEN)] == 'descr-preserved-pre-existing'
        # flag on: only the autogen one replaced
        _apply(work, '--replace-autogen-descr')
        acts = _audit_actions(work, deck)
        assert acts[_h(PNG_RED)] == 'descr-replaced-autogen'
        assert acts[_h(PNG_GREEN)] == 'descr-preserved-pre-existing'
        _, xml = _saved_slide_xml(work, deck)
        assert 'descr="whiteboard sketch."' in xml
        assert 'descr="Real alt text."' in xml
        print('  --replace-autogen-descr OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_descr_only_mode():
    tmp = tempfile.mkdtemp(prefix='descr_t4_')
    try:
        deck_path = os.path.join(tmp, 'deckC.pptx')
        _build_deck(deck_path, [(PNG_RED, None)])
        caps = {_h(PNG_RED): 'Photo: coffee mug on red background.'}
        work, deck = _make_work(tmp, deck_path, caps)
        subprocess.run([sys.executable, os.path.join(SCRIPTS, 'apply_captions.py'),
                        work, '--quick', '--descr-only'],
                       capture_output=True, text=True, check=True)
        out, xml = _saved_slide_xml(work, deck)
        assert 'descr="coffee mug on red background."' in xml
        assert 'captioner_caption_' not in xml, 'descr-only must add NO text boxes'
        print('  --descr-only OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_preexisting_decorative_mark_respected():
    """A picture the instructor already marked decorative in PowerPoint must
    NOT receive authored descr text — the mark is itself a decorative signal."""
    tmp = tempfile.mkdtemp(prefix='descr_t6_')
    try:
        deck_path = os.path.join(tmp, 'deckE.pptx')
        _build_deck(deck_path, [(PNG_RED, None)])
        prs = Presentation(deck_path)
        (el, cnvpr, pid) = iter_live_pic_cnvprs(prs.slides[0])[0]
        mark_decorative(cnvpr)
        prs.save(deck_path)
        caps = {_h(PNG_RED): 'Photo: coffee mug on red background.'}
        work, deck = _make_work(tmp, deck_path, caps)
        _apply(work)
        acts = _audit_actions(work, deck)
        assert acts[_h(PNG_RED)] == 'descr-decorative-unchanged', acts
        _, xml = _saved_slide_xml(work, deck)
        assert 'descr="coffee mug' not in xml
        assert DECOR_EXT_URI in xml
        print('  pre-existing decorative mark respected OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_verify_descr_gate():
    tmp = tempfile.mkdtemp(prefix='descr_t5_')
    try:
        deck_path = os.path.join(tmp, 'deckD.pptx')
        _build_deck(deck_path, [(PNG_RED, None), (PNG_GREEN, 'Hand authored text'),
                                (PNG_BLUE, None)])
        caps = {_h(PNG_RED): 'Photo: coffee mug on red background.',
                _h(PNG_GREEN): 'Photo: green square artwork.',
                _h(PNG_BLUE): '[decorative]'}
        work, deck = _make_work(tmp, deck_path, caps)
        _apply(work)
        r = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'verify.py'),
                            work, '--descr'], capture_output=True, text=True)
        assert 'authored=1' in r.stdout and 'preserved-pre-existing' in r.stdout
        assert 'MISSING=0' in r.stdout
        assert r.returncode == 0, f'verify rc={r.returncode}\n{r.stdout}'
        # remove one descr -> gate must exit 6
        out = os.path.join(work, 'captioned_decks', f'{deck}_captioned.pptx')
        prs = Presentation(out)
        for (el, cnvpr, pid) in iter_live_pic_cnvprs(prs.slides[0]):
            if (cnvpr.get('descr') or '') == 'coffee mug on red background.':
                del cnvpr.attrib['descr']
        prs.save(out)
        r2 = subprocess.run([sys.executable, os.path.join(SCRIPTS, 'verify.py'),
                             work, '--descr'], capture_output=True, text=True)
        assert r2.returncode == 6, f'expected exit 6, got {r2.returncode}\n{r2.stdout}'
        print('  verify --descr gate (0 clean / 6 missing) OK')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    print('descr-writer tests:')
    test_decide_descr_matrix()
    test_derive_alt_text()
    test_is_autogen_descr()
    test_mark_decorative_roundtrip()
    test_apply_end_to_end()
    test_replace_autogen_flag()
    test_descr_only_mode()
    test_preexisting_decorative_mark_respected()
    test_verify_descr_gate()
    print('ALL DESCR TESTS PASSED')
